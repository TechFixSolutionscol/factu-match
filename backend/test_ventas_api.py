"""Endpoints /api/ventas/… con Odoo y Claude falsos (sin red)."""

import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import ventas.router as rutas
from ventas import almacen
from ventas.odoo_ventas import OdooSoloLectura


@pytest.fixture
def cliente(ventas, monkeypatch):
    almacen.limpiar()
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    app = FastAPI()
    app.include_router(rutas.router)
    filas = [ventas.fila("FE", 1, 1000), ventas.fila("FE", 2, 2000), ventas.fila("FE", 3, 3000)]
    odoo = ventas.OdooFalso(movimientos=[
        ventas.movimiento(1, "FE/2026/00001", 1000, cufe_valor=filas[0]["CUFE/CUDE"]),
        ventas.movimiento(2, "FE/2026/00002", 2500, cufe_valor=filas[1]["CUFE/CUDE"]),
    ])
    monkeypatch.setattr(rutas, "_lector", lambda credenciales, segundos=120.0: OdooSoloLectura(odoo))
    c = TestClient(app)
    c.contenedor = ventas.contenedor(filas)
    c.ventas = ventas
    yield c
    almacen.limpiar()


def _conciliar(c, **form):
    datos = {"credentials": "token", "desde": "2026-10-01", "hasta": "2026-10-15", **form}
    return c.post("/api/ventas/conciliar", files={"dian": ("contenedor.xlsx", c.contenedor)}, data=datos)


def test_analizar_contenedor(cliente):
    r = cliente.post("/api/ventas/analizar-contenedor", files={"dian": ("contenedor.xlsx", cliente.contenedor)})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["documentos_incluidos"] == 3 and d["bloqueante"] is False
    assert d["periodo_sugerido"] == {"desde": "2026-10-05", "hasta": "2026-10-05"}
    assert len(d["mapeo_columnas"]) == 32 and all(c["reconocida"] for c in d["mapeo_columnas"])


@pytest.mark.parametrize("nombre, contenido, estado", [
    ("contenedor.csv", b"a,b", 400),
    ("contenedor.xlsx", b"", 400),
    ("contenedor.xlsx", b"%PDF-1.7 no es excel", 422),
])
def test_archivos_invalidos(cliente, nombre, contenido, estado):
    r = cliente.post("/api/ventas/analizar-contenedor", files={"dian": (nombre, contenido)})
    assert r.status_code == estado


def test_conciliar_documento_y_reporte(cliente):
    r = _conciliar(cliente)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["descomposicion"]["diferencia_neta"] == "2500.00" and d["descomposicion"]["cuadra"] is True
    assert d["resumen"]["por_estado"]["missing_in_odoo"] == 1
    assert d["resumen"]["por_estado"]["amount_mismatch"] == 1

    doc = cliente.get("/api/ventas/documento", params={"run_id": d["run_id"], "ref": "dian:FE-2"})
    assert doc.status_code == 200 and doc.json()["estado"] == "amount_mismatch"
    assert doc.json()["validaciones"]
    assert cliente.get("/api/ventas/documento", params={"run_id": d["run_id"], "ref": "dian:XX-1"}).status_code == 404

    rep = cliente.get("/api/ventas/reporte", params={"run_id": d["run_id"]})
    assert rep.status_code == 200
    assert "conciliacion_ventas_2026-10-01_2026-10-15.xlsx" in rep.headers["content-disposition"]
    assert "Hallazgos IA" in load_workbook(io.BytesIO(rep.content)).sheetnames


def test_parametros_invalidos(cliente):
    assert _conciliar(cliente, tolerancia="NaN").status_code == 422
    assert _conciliar(cliente, tolerancia="sNaN").status_code == 422
    assert _conciliar(cliente, tolerancia="-1").status_code == 422
    assert _conciliar(cliente, desde="2026-10-20", hasta="2026-10-01").status_code == 422
    assert _conciliar(cliente, desde="2026-01-01", hasta="2026-12-31").status_code == 422
    assert _conciliar(cliente, tolerancia="abc").status_code == 422
    assert _conciliar(cliente, incluir_move="false", incluir_pos="false").status_code == 422


def test_contenedor_con_varios_nit_emisores_se_rechaza(cliente):
    v = cliente.ventas
    cliente.contenedor = v.contenedor([v.fila("FE", 1, 10), v.fila("FE", 2, 10, nit_emisor="800000002")])
    r = _conciliar(cliente)
    assert r.status_code == 422 and "NIT" in r.json()["detail"]


def test_compania_distinta_en_odoo(cliente, monkeypatch):
    v = cliente.ventas
    otra = v.OdooFalso(companias=[{"id": 1, "name": "Otra", "vat": "800000002"}])
    monkeypatch.setattr(rutas, "_lector", lambda credenciales, segundos=120.0: OdooSoloLectura(otra))
    r = _conciliar(cliente)
    assert r.status_code == 422


def test_credenciales_invalidas(cliente, monkeypatch):
    monkeypatch.undo()                                        # usa el _lector real
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    r = _conciliar(cliente, credentials="no-es-un-token-fernet")
    assert r.status_code == 400 and "Credenciales de Odoo" in r.json()["detail"]


def test_conciliacion_expirada(cliente):
    assert cliente.get("/api/ventas/reporte", params={"run_id": "noexiste"}).status_code == 404
    assert cliente.post("/api/ventas/investigar", data={"run_id": "noexiste"}).status_code == 404


def test_investigar_sin_clave_responde_no_disponible(cliente):
    run_id = _conciliar(cliente).json()["run_id"]
    r = cliente.post("/api/ventas/investigar", data={"run_id": run_id})
    assert r.status_code == 200
    assert r.json()["disponible"] is False and r.json()["codigo"] == "no_configurada"
    assert cliente.get("/api/ventas/reporte", params={"run_id": run_id}).status_code == 200


def test_investigar_con_falla_del_proveedor_mantiene_el_reporte(cliente, monkeypatch):
    import anthropic
    import httpx2

    from ventas.agente import proveedor

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-prueba")

    class Caido:
        @property
        def messages(self):
            return self

        def create(self, **_):
            req = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
            raise anthropic.RateLimitError("rate", response=httpx2.Response(429, request=req), body=None)

    monkeypatch.setattr(proveedor, "_crear_cliente", lambda config: Caido())
    run_id = _conciliar(cliente).json()["run_id"]
    r = cliente.post("/api/ventas/investigar", data={"run_id": run_id, "foco_tipo": "componente", "foco_valor": "missing_in_odoo"})
    assert r.status_code == 200 and r.json()["codigo"] == "cuota"
    rep = cliente.get("/api/ventas/reporte", params={"run_id": run_id})
    hoja = load_workbook(io.BytesIO(rep.content))["Hallazgos IA"]
    assert "Análisis IA no disponible" in hoja["A3"].value


def test_investigacion_en_curso_no_se_duplica(cliente):
    run_id = _conciliar(cliente).json()["run_id"]
    rutas._investigaciones_en_curso.add(run_id)
    try:
        assert cliente.post("/api/ventas/investigar", data={"run_id": run_id}).status_code == 409
        assert cliente.post("/api/ventas/investigar", data={"run_id": f" {run_id} "}).status_code == 409
    finally:
        rutas._investigaciones_en_curso.discard(run_id)


def test_estado_ia_nunca_devuelve_la_clave(cliente, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-no-debe-salir")
    r = cliente.get("/api/ventas/estado-ia")
    assert r.status_code == 200 and r.json()["configurada"] is True
    assert "sk-ant-no-debe-salir" not in r.text
    assert r.json()["modelo"] == "claude-haiku-5-5"


def test_main_registra_ventas_sin_tocar_compras():
    import main
    rutas_app = {r.path for r in main.app.routes}
    assert {"/api/ventas/conciliar", "/api/ventas/investigar", "/api/ventas/reporte"} <= rutas_app
    assert {"/comparar", "/descargar-reporte", "/odoo/comparar", "/odoo/descargar-reporte",
            "/api/purchase/pending", "/api/purchase/create-oc"} <= rutas_app


def test_archivo_demasiado_grande(cliente, monkeypatch):
    monkeypatch.setattr(rutas, "MAX_ARCHIVO", 1024)
    r = cliente.post("/api/ventas/analizar-contenedor", files={"dian": ("contenedor.xlsx", cliente.contenedor)})
    assert r.status_code == 413


def test_permiso_insuficiente_en_odoo(cliente, monkeypatch):
    import xmlrpc.client
    v = cliente.ventas

    class SinPermiso(v.OdooFalso):
        def __call__(self, modelo, metodo, args, kwargs):
            if modelo == "account.move" and metodo == "search_read":
                raise xmlrpc.client.Fault(4, "AccessError: no puede leer facturas")
            return super().__call__(modelo, metodo, args, kwargs)

    monkeypatch.setattr(rutas, "_lector", lambda credenciales, segundos=120.0: OdooSoloLectura(SinPermiso()))
    r = _conciliar(cliente)
    assert r.status_code == 422 and "permiso" in r.json()["detail"]


def test_investigar_sigue_si_odoo_en_vivo_falla(cliente, monkeypatch):
    from types import SimpleNamespace as NS

    from ventas.agente import proveedor

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-prueba")
    run_id = _conciliar(cliente).json()["run_id"]

    def caido(credenciales, segundos=120.0):
        raise RuntimeError("socket roto")

    class Final:
        @property
        def messages(self):
            return self

        def create(self, **kwargs):
            vacio = {"resumen_general": "Sin hallazgos.", "hallazgos": [], "datos_faltantes": [], "siguiente_conciliacion": ""}
            texto = "Listo." if kwargs["tool_choice"] == {"type": "auto"} else __import__("json").dumps(vacio)
            return NS(content=[NS(type="text", text=texto)], stop_reason="end_turn", usage=None)

    monkeypatch.setattr(rutas, "_lector", caido)
    monkeypatch.setattr(proveedor, "_crear_cliente", lambda config: Final())
    r = cliente.post("/api/ventas/investigar", data={"run_id": run_id, "credentials": "token"})
    assert r.status_code == 200 and r.json()["disponible"] is True
    assert "Sin consulta en vivo" in r.json()["aviso_odoo"]
