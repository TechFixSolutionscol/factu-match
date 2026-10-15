"""
Reporte Excel de la conciliación de ventas.

Independiente del reporte de compras (comparador.generar_excel_reporte no se toca).
Se genera siempre con el resultado determinístico; si hubo análisis de Claude,
se agrega en su propia hoja, marcado como interpretación y no como hecho contable.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .agente.esquema import ETIQUETAS_ESTADO_CAUSA, ETIQUETAS_NIVEL, ETIQUETAS_RIESGO
from .motor import ETIQUETAS_ESTADO, ResultadoConciliacion
from .normalizacion import a_decimal, cufe_corto, fecha_iso

AZUL = "1A3A5C"
AZUL_MEDIO = "185FA5"
ROJO = "C00000"
VIOLETA = "5B3B8C"
FORMATO_DINERO = '#,##0.00;[Red]-#,##0.00'
_BORDE = Border(*(Side(style="thin", color="D0D7E2"),) * 4)

ESTADOS_PENDIENTES = ("missing_in_odoo", "missing_in_dian", "amount_mismatch", "tax_mismatch",
                      "identity_mismatch", "possible_duplicate", "ambiguous", "needs_review", "scope_difference")


def _titulo(ws, texto: str, columnas: int, color: str = AZUL) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columnas)
    celda = ws.cell(row=1, column=1, value=texto)
    celda.font = Font(bold=True, color="FFFFFF", size=12)
    celda.fill = PatternFill("solid", fgColor=color)
    celda.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22


def _encabezado(ws, fila: int, nombres: list[str], color: str = AZUL_MEDIO) -> None:
    for j, nombre in enumerate(nombres, start=1):
        c = ws.cell(row=fila, column=j, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDE


_ARRIBA = Alignment(vertical="top")
_ARRIBA_AJUSTADO = Alignment(vertical="top", wrap_text=True)


def _fila(ws, fila: int, valores: list, dinero_cols: tuple[int, ...] = (), ligera: bool = False) -> None:
    """Escribe una fila. `ligera` omite bordes y alineación: en hojas de miles de filas el estilo
    por celda es lo que más tarda (openpyxl) y no aporta a la lectura."""
    for j, valor in enumerate(valores, start=1):
        if j in dinero_cols and valor not in (None, ""):
            valor = a_decimal(valor) if not isinstance(valor, Decimal) else valor
        c = ws.cell(row=fila, column=j, value=valor)
        if not ligera:
            c.border = _BORDE
            c.alignment = _ARRIBA_AJUSTADO if isinstance(valor, str) and len(valor) > 40 else _ARRIBA
        if j in dinero_cols and valor not in (None, ""):
            c.number_format = FORMATO_DINERO


def _anchos(ws, anchos: list[int]) -> None:
    for j, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = ancho


def generar_reporte(resultado: ResultadoConciliacion) -> bytes:
    datos = resultado.a_dict(incluir_filas=False)
    wb = Workbook()

    # 1. Resumen ejecutivo
    ws = wb.active
    ws.title = "Resumen"
    _titulo(ws, "CONCILIACIÓN DE VENTAS DIAN VS ODOO — RESUMEN EJECUTIVO", 4)
    r = datos["resumen"]
    desc = datos["descomposicion"]
    neto = next(m for m in datos["metricas"] if m["codigo"] == "total_neto")
    filas = [
        ("Empresa", f"{datos['odoo']['compania']['nombre']} · NIT {datos['odoo']['compania']['nit']}"),
        ("Período", f"{datos['parametros']['desde']} → {datos['parametros']['hasta']}"),
        ("Generado (UTC)", datos["generado_en"]),
        ("Contenedor DIAN", f"{datos['contenedor']['archivo']} · hoja {datos['contenedor']['hoja']}"),
        ("Documentos DIAN elegibles", r["documentos_dian_elegibles"]),
        ("Documentos Odoo canónicos", r["documentos_odoo_canonicos"]),
        ("Emparejados", r["emparejados"]),
        ("Total neto DIAN", a_decimal(neto["dian"])),
        ("Total neto Odoo", a_decimal(neto["odoo"])),
        ("Diferencia neta (DIAN − Odoo)", a_decimal(desc["diferencia_neta"])),
        ("Explicado por componentes", a_decimal(desc["explicado"])),
        ("Residuo no explicado", a_decimal(desc["residuo"])),
    ]
    for i, (k, v) in enumerate(filas, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, Decimal):
            c.number_format = FORMATO_DINERO
    fila = len(filas) + 4
    _encabezado(ws, fila, ["Estado", "Código", "Documentos"])
    for codigo, n in r["por_estado"].items():
        fila += 1
        _fila(ws, fila, [ETIQUETAS_ESTADO.get(codigo, codigo), codigo, n])
    fila += 2
    ws.cell(row=fila, column=1, value="Análisis de Claude").font = Font(bold=True)
    ws.cell(row=fila, column=2, value="Ver hoja «Hallazgos IA»" if resultado.hallazgos_ia
            else "No ejecutado o no disponible: el resultado de este reporte es 100 % determinístico")
    _anchos(ws, [34, 60, 14, 14])

    # 2. Parámetros y alcance
    ws = wb.create_sheet("Parámetros")
    _titulo(ws, "PARÁMETROS DE EJECUCIÓN Y ALCANCE COMPARADO", 3)
    p = datos["parametros"]
    filas = [("Desde", p["desde"]), ("Hasta", p["hasta"]), ("Criterio de fecha", p["criterio_fecha"]),
             ("Tolerancia por documento", p["tolerancia_documento"]), ("Diferencia global", p["diferencia_global"]),
             ("Fuente account.move", "sí" if p["fuentes"]["account.move"] else "no"),
             ("Fuente pos.order", "sí" if p["fuentes"]["pos.order"] else "no"),
             ("Corte del contenedor", datos["contenedor"]["corte_recepcion"] or "—")]
    for i, (k, v) in enumerate(filas, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    fila = len(filas) + 4
    _encabezado(ws, fila, ["Criterio", "¿Coincide?", "Detalle"])
    for a in datos["alcance"]:
        fila += 1
        _fila(ws, fila, [a["etiqueta"], "sí" if a["ok"] else "NO", a["detalle"]])
    fila += 2
    _encabezado(ws, fila, ["Tipo de documento (valor exacto)", "Filas", "Tratamiento"])
    for t in datos["contenedor"]["tipos"]:
        fila += 1
        _fila(ws, fila, [t["tipo_documento"], t["filas"], t["tratamiento"]])
    _anchos(ws, [38, 14, 90])

    # 3. Global
    ws = wb.create_sheet("Global")
    _titulo(ws, "CONCILIACIÓN GLOBAL DE TOTALES", 6)
    _encabezado(ws, 3, ["Métrica", "DIAN", "Odoo", "Diferencia", "Dif. %", "Nota"])
    for i, m in enumerate(datos["metricas"], start=4):
        es_cantidad = m["codigo"].startswith("cantidad")
        _fila(ws, i, [m["etiqueta"], m["dian"] if not es_cantidad else int(m["dian"]),
                      m["odoo"] if (m["odoo"] is None or not es_cantidad) else int(m["odoo"]),
                      m["diferencia"] if (m["diferencia"] is None or not es_cantidad) else int(m["diferencia"]),
                      m["diferencia_porcentual"], m["nota"]],
              dinero_cols=() if es_cantidad else (2, 3, 4))
    _anchos(ws, [40, 18, 18, 18, 10, 50])

    # 4. Descomposición
    ws = wb.create_sheet("Descomposición")
    _titulo(ws, "DESCOMPOSICIÓN DE LA DIFERENCIA GLOBAL", 6)
    _encabezado(ws, 3, ["Componente", "Código", "Signo", "Documentos", "Monto (DIAN − Odoo)", "Participación %"])
    fila = 3
    for c in desc["componentes"]:
        fila += 1
        _fila(ws, fila, [c["etiqueta"], c["codigo"], c["signo"], c["documentos"], c["monto"], c["participacion"]], dinero_cols=(5,))
    for etiqueta, valor in (("Diferencia neta", desc["diferencia_neta"]), ("Explicado", desc["explicado"]),
                            ("Residuo no explicado", desc["residuo"])):
        fila += 1
        _fila(ws, fila, [etiqueta, "", "", "", valor, ""], dinero_cols=(5,))
        ws.cell(row=fila, column=1).font = Font(bold=True)
    fila += 2
    ws.cell(row=fila, column=1, value=desc["formula"])
    _anchos(ws, [58, 26, 8, 12, 22, 14])

    # 5. Prefijos
    ws = wb.create_sheet("Prefijos")
    _titulo(ws, "POR PREFIJO (PUNTO DE VENTA O CANAL)", 9)
    _encabezado(ws, 3, ["Prefijo", "Origen en Odoo", "Facturas DIAN", "NC DIAN", "Neto DIAN", "Neto Odoo",
                        "Diferencia", "Folios", "Ausentes en secuencia"])
    for i, pfx in enumerate(datos["prefijos"], start=4):
        rango = f"{pfx['folio_minimo']} → {pfx['folio_maximo']}" if pfx["folio_minimo"] is not None else "—"
        _fila(ws, i, [pfx["prefijo"], pfx["origen"] or "—", pfx["facturas_dian"], pfx["notas_credito_dian"],
                      pfx["neto_dian"], pfx["neto_odoo"], pfx["diferencia"], rango, pfx["ausentes_en_secuencia"]],
              dinero_cols=(5, 6, 7))
    _anchos(ws, [10, 28, 14, 10, 18, 18, 18, 18, 12])

    # 6. Secuencias
    ws = wb.create_sheet("Secuencias")
    _titulo(ws, "FOLIOS AUSENTES EN EL CONTENEDOR (POR PREFIJO)", 6)
    _encabezado(ws, 3, ["Prefijo", "Bloque", "Cantidad", "Folio anterior", "Folio siguiente", "Fechas de recepción vecinas"])
    fila = 3
    for s in datos["secuencias"]:
        for b in s["bloques"]:
            fila += 1
            ant, sig = b["anterior"] or {}, b["siguiente"] or {}
            _fila(ws, fila, [s["prefijo"], f"{b['desde']}–{b['hasta']}", b["cantidad"], ant.get("numero", "—"),
                             sig.get("numero", "—"), f"{ant.get('fecha_recepcion') or '—'} / {sig.get('fecha_recepcion') or '—'}"])
    _anchos(ws, [10, 16, 10, 16, 16, 44])

    # 7. Detalle por documento
    ws = wb.create_sheet("Detalle")
    _titulo(ws, "DETALLE POR DOCUMENTO", 16)
    cab = ["Estado", "Código estado", "Método", "Número", "Tipo", "Fecha", "NIT receptor", "Receptor",
           "Total DIAN", "IVA DIAN", "INC DIAN", "Registro Odoo", "Total Odoo", "Dif. total", "Dif. impuestos", "Evidencia"]
    _encabezado(ws, 3, cab)
    for i, f in enumerate(resultado.filas, start=4):
        d, o = f.dian, f.odoo
        _fila(ws, i, [
            ETIQUETAS_ESTADO.get(f.estado, f.estado), f.estado, f.metodo,
            d.numero_legible if d else (o.numero_legible if o else ""),
            d.tipo_documento if d else o.tipo,
            fecha_iso(f.fecha), d.nit_receptor if d else (o.partner_vat if o else ""),
            d.nombre_receptor if d else (o.partner_nombre if o else ""),
            d.total if d else None, d.iva if d else None, d.inc if d else None,
            o.nombre if o else ", ".join(c.nombre for c in f.candidatos), o.total if o else None,
            f.diferencia_total, f.diferencia_impuestos, " · ".join(f.evidencia),
        ], dinero_cols=(9, 10, 11, 13, 14, 15), ligera=True)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(cab))}{max(4, len(resultado.filas) + 3)}"
    _anchos(ws, [20, 18, 14, 14, 24, 11, 14, 32, 15, 12, 12, 22, 15, 13, 13, 70])

    # 8. Evidencia de Odoo
    ws = wb.create_sheet("Evidencia Odoo")
    _titulo(ws, "REGISTROS DE ODOO USADOS COMO EVIDENCIA (SOLO LECTURA)", 11)
    _encabezado(ws, 3, ["Referencia", "Modelo", "ID", "Nombre", "Números", "CUFE", "Fecha", "Estado",
                        "Estado electrónico", "Diario / POS", "Alcance"])
    fila = 3
    for doc in resultado.conjunto.documentos + resultado.conjunto.fuera_de_alcance:
        fila += 1
        _fila(ws, fila, [doc.ref, doc.modelo, doc.registro_id, doc.nombre, ", ".join(n.legible for n in doc.numeros),
                         cufe_corto(doc.cufe), fecha_iso(doc.fecha), doc.estado, doc.estado_electronico, doc.diario,
                         "en alcance" if doc.en_alcance else doc.motivo_alcance], ligera=True)
    ws.freeze_panes = "A4"
    _anchos(ws, [28, 13, 9, 24, 20, 20, 11, 10, 16, 20, 50])

    # 9. Hallazgos de Claude (interpretación, no hecho contable)
    ws = wb.create_sheet("Hallazgos IA")
    _titulo(ws, "HALLAZGOS DE CLAUDE — INTERPRETACIÓN, NO HECHO CONTABLE", 4, color=VIOLETA)
    h = resultado.hallazgos_ia
    if not h or not h.get("disponible"):
        ws.cell(row=3, column=1, value=(h or {}).get("mensaje") or "No se ejecutó el análisis de Claude para esta conciliación.")
        ws.cell(row=4, column=1, value="La conciliación de este reporte es 100 % determinística y no depende de la IA.")
    else:
        r = h["resultado"]
        fila = 3
        for k, v in (("Modelo", h.get("modelo")), ("Generado (UTC)", h.get("generado_en")),
                     ("Foco", " ".join(str(x) for x in (h.get("foco") or {}).values())),
                     ("Aviso", h.get("mensaje")), ("Resumen", r.get("resumen_general")),
                     ("Siguiente conciliación", r.get("siguiente_conciliacion")),
                     ("Datos faltantes", " · ".join(r.get("datos_faltantes") or []) or "—"),
                     ("Escala de confianza", "Alta / Media / Baja: cualitativa, no es una probabilidad calibrada")):
            ws.cell(row=fila, column=1, value=k).font = Font(bold=True)
            ws.cell(row=fila, column=2, value=v)
            fila += 1
        for n, hz in enumerate(r["hallazgos"], start=1):
            fila += 1
            ws.cell(row=fila, column=1, value=f"Hallazgo {n}: {hz['summary']}").font = Font(bold=True)
            fila += 1
            _fila(ws, fila, ["Tipo · estado de la causa",
                             f"{hz.get('etiqueta_tipo', hz['finding_type'])} · {ETIQUETAS_ESTADO_CAUSA.get(hz['cause_status'], hz['cause_status'])}",
                             "Riesgo · confianza",
                             f"{ETIQUETAS_RIESGO.get(hz['risk_level'], hz['risk_level'])} · {ETIQUETAS_NIVEL.get(hz['confidence'], hz['confidence'])}"])
            impacto = hz.get("financial_impact") or {}
            citados = impacto.get("documentos_citados") or {}
            componente = impacto.get("componente") or {}
            fila += 1
            _fila(ws, fila, ["Impacto (calculado por el backend)",
                             a_decimal(citados["monto"]) if citados.get("monto") else None,
                             f"{citados.get('documentos', 0)} documentos citados",
                             f"Componente {componente.get('etiqueta')}: {componente.get('monto')} ({componente.get('documentos')} docs)"
                             if componente else ""], dinero_cols=(2,))
            for hecho in hz["facts"]:
                fila += 1
                _fila(ws, fila, ["HECHO (evidencia del backend)", hecho["texto"], "Evidencia", ", ".join(hecho["evidencia"])])
            for causa in hz["probable_causes"]:
                fila += 1
                _fila(ws, fila, ["HIPÓTESIS SIN CONFIRMAR", causa["descripcion"],
                                 f"Confianza {ETIQUETAS_NIVEL.get(causa['confianza'], causa['confianza'])}",
                                 "Cómo confirmarla: " + causa["como_confirmar"]])
            for accion in hz["recommended_actions"]:
                fila += 1
                _fila(ws, fila, ["ACCIÓN PENDIENTE DEL USUARIO", accion["accion"], accion["registro_afectado"],
                                 "Verificación: " + accion["verificacion"]])
            for paso in hz["verification_steps"]:
                fila += 1
                _fila(ws, fila, ["VERIFICACIÓN", paso, "", ""])
            for aviso in hz.get("advertencias", []):
                fila += 1
                _fila(ws, fila, ["VALIDACIÓN DEL BACKEND", aviso, "", ""])
            if hz.get("requires_human_review"):
                fila += 1
                _fila(ws, fila, ["REQUIERE REVISIÓN HUMANA", "Sí", "", ""])
        descartados = (h.get("validacion") or {}).get("descartados") or []
        if descartados:
            fila += 2
            ws.cell(row=fila, column=1, value="Hallazgos descartados por falta de evidencia").font = Font(bold=True)
            for d in descartados:
                fila += 1
                _fila(ws, fila, [f"Hallazgo {d['hallazgo']}", d["resumen"], "", d["motivo"]])
    _anchos(ws, [32, 80, 24, 60])

    # 10. Pendientes de revisión
    ws = wb.create_sheet("Pendientes")
    _titulo(ws, "CASOS PENDIENTES DE REVISIÓN (ORDENADOS POR IMPACTO)", 6, color=ROJO)
    _encabezado(ws, 3, ["Estado", "Número", "Fecha", "Aporte a la diferencia", "Registro Odoo", "Evidencia"])
    pendientes = sorted((f for f in resultado.filas if f.estado in ESTADOS_PENDIENTES),
                        key=lambda f: abs(f.contribucion), reverse=True)
    for i, f in enumerate(pendientes, start=4):
        d, o = f.dian, f.odoo
        _fila(ws, i, [ETIQUETAS_ESTADO.get(f.estado, f.estado), d.numero_legible if d else (o.numero_legible if o else ""),
                      fecha_iso(f.fecha), f.contribucion, o.nombre if o else ", ".join(c.nombre for c in f.candidatos),
                      " · ".join(f.evidencia)], dinero_cols=(4,))
    ws.freeze_panes = "A4"
    _anchos(ws, [22, 14, 11, 20, 24, 90])

    _neutralizar_formulas(wb)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _neutralizar_formulas(wb: Workbook) -> None:
    """
    El reporte no usa fórmulas. Un texto que empiece por «=» (nombre de cliente, nombre de un
    registro de Odoo o texto de Claude) se guardaría como fórmula y Excel la ejecutaría
    (inyección de fórmulas): se fuerza a texto.
    """
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for celda in fila:
                if celda.data_type == "f":
                    celda.data_type = "s"
