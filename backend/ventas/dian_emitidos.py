"""
Lector del contenedor DIAN de documentos emitidos.

Usa los nombres exactos de columna del Excel que exporta el portal de la DIAN
(hoja «Rp_Doc_…»): Tipo de documento, CUFE/CUDE, Folio, Prefijo, Divisa,
Forma de Pago, Medio de Pago, Fecha Emisión, Fecha Recepción, NIT Emisor,
Nombre Emisor, NIT Receptor, Nombre Receptor, IVA, ICA, IC, INC, Timbre,
INC Bolsas, IN Carbono, IN Combustibles, IC Datos, ICL, INPP, IBUA, ICUI,
Rete IVA, Rete Renta, Rete ICA, Total, Estado, Grupo.

La comparación de encabezados y de valores de «Tipo de documento» ignora
mayúsculas, tildes y espacios dobles; el valor exacto se conserva para mostrarlo.

No modifica ni reutiliza la lógica de compras: solo toma prestada la detección
de formato de archivo de comparador.py.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Optional

from comparador import _detect_excel_format

from .normalizacion import (
    CERO,
    NumeroDocumento,
    a_decimal_seguro,
    cufe_corto,
    dinero,
    es_consumidor_final,
    fecha_iso,
    importe_ambiguo,
    normalizar_cufe,
    normalizar_nit,
    normalizar_prefijo,
    numero_desde_dian,
    parsear_fecha,
    parsear_fecha_hora,
    texto_clave,
    texto_limpio,
)

# ──────────────────────────────────────────────
# COLUMNAS (nombres exactos del portal DIAN)
# ──────────────────────────────────────────────

COLUMNAS_REQUERIDAS = (
    "Tipo de documento", "CUFE/CUDE", "Folio", "Prefijo", "Fecha Emisión",
    "NIT Emisor", "NIT Receptor", "IVA", "INC", "Total", "Estado", "Grupo",
)
COLUMNAS_CONTEXTO = (
    "Divisa", "Forma de Pago", "Medio de Pago", "Fecha Recepción",
    "Nombre Emisor", "Nombre Receptor",
)
OTROS_IMPUESTOS = (
    "ICA", "IC", "Timbre", "INC Bolsas", "IN Carbono", "IN Combustibles",
    "IC Datos", "ICL", "INPP", "IBUA", "ICUI",
)
RETENCIONES = ("Rete IVA", "Rete Renta", "Rete ICA")
TODAS_LAS_COLUMNAS = COLUMNAS_REQUERIDAS + COLUMNAS_CONTEXTO + OTROS_IMPUESTOS + RETENCIONES

GRUPO_EMITIDO = "emitido"
MAX_FILAS = 100_000
MAX_COLUMNAS = 100                               # el contenedor tiene 32; las demás no se usan
MAX_XLSX_DESCOMPRIMIDO = 200 * 1024 * 1024       # un contenedor real de 10 MB descomprime en ~70 MB
_FILA_XML = re.compile(rb"<(?:[A-Za-z_][\w.-]{0,20}:)?row[\s/>]")
ESTADOS_APROBADOS = {"aprobado", "aprobado con notificacion"}


# ──────────────────────────────────────────────
# REGLAS POR TIPO DE DOCUMENTO
# ──────────────────────────────────────────────

@dataclass(frozen=True)
class ReglaTipo:
    categoria: str  # factura | nota_credito | nota_debito | excluido
    signo: int
    incluir: bool
    regla: str


_VENTA = ReglaTipo("factura", 1, True, "Venta: suma al total bruto")
_NC = ReglaTipo("nota_credito", -1, True, "Nota crédito: se resta del neto, separada del bruto")
_ND = ReglaTipo("nota_debito", 1, True, "Nota débito: suma al neto, separada del bruto")
_SOPORTE = ReglaTipo("excluido", 0, False, "Documento soporte: es compra, no venta (decisión del 15-10-2026)")

REGLAS_TIPO: dict[str, ReglaTipo] = {
    "factura electronica": _VENTA,
    "factura electronica de venta": _VENTA,
    "documento equivalente pos": ReglaTipo("factura", 1, True, "Documento equivalente POS: suma como venta"),
    "documento equivalente electronico pos": ReglaTipo("factura", 1, True, "Documento equivalente POS: suma como venta"),
    "nota de credito electronica": _NC,
    "nota credito electronica": _NC,
    "nota de debito electronica": _ND,
    "nota debito electronica": _ND,
    "application response": ReglaTipo("excluido", 0, False, "Evento o acuse (Application response): no es venta"),
    "nomina individual": ReglaTipo("excluido", 0, False, "Nómina electrónica: no es venta"),
    "nomina individual de ajuste": ReglaTipo("excluido", 0, False, "Nómina electrónica de ajuste: no es venta"),
    "documento soporte con no obligados": _SOPORTE,
    "nota de ajuste del documento soporte": _SOPORTE,
}


def regla_para_tipo(tipo: str) -> Optional[ReglaTipo]:
    return REGLAS_TIPO.get(texto_clave(tipo))


# ──────────────────────────────────────────────
# MODELOS
# ──────────────────────────────────────────────

@dataclass
class DocumentoDian:
    fila: int
    tipo_documento: str
    categoria: str
    signo: int
    incluido: bool
    motivo: str
    cufe: str
    prefijo: str
    folio: str
    numero: Optional[NumeroDocumento]
    fecha_emision: Optional[date]
    fecha_recepcion: Optional[datetime]
    nit_emisor: str
    nombre_emisor: str
    nit_receptor: str
    nombre_receptor: str
    divisa: str
    forma_pago: str
    medio_pago: str
    iva: Decimal
    inc: Decimal
    otros_impuestos: dict[str, Decimal]
    retenciones: dict[str, Decimal]
    total: Decimal
    estado_dian: str
    grupo: str
    advertencias: list[str] = field(default_factory=list)

    @property
    def ref(self) -> str:
        return f"dian:{self.numero.legible}" if self.numero else f"dian:fila{self.fila}"

    @property
    def numero_legible(self) -> str:
        return self.numero.legible if self.numero else (self.folio or f"fila {self.fila}")

    @property
    def total_otros_impuestos(self) -> Decimal:
        return sum(self.otros_impuestos.values(), CERO)

    @property
    def impuestos(self) -> Decimal:
        return self.iva + self.inc + self.total_otros_impuestos

    @property
    def base_derivada(self) -> Decimal:
        """El contenedor no trae subtotal: base = Total − IVA − INC − otros impuestos."""
        return self.total - self.impuestos

    @property
    def consumidor_final(self) -> bool:
        return es_consumidor_final(self.nit_receptor)

    @property
    def total_firmado(self) -> Decimal:
        return self.total * self.signo

    def a_dict(self) -> dict:
        return {
            "ref": self.ref,
            "fila": self.fila,
            "tipo_documento": self.tipo_documento,
            "categoria": self.categoria,
            "incluido": self.incluido,
            "motivo": self.motivo,
            "numero": self.numero_legible,
            "prefijo": self.prefijo,
            "folio": self.folio,
            "cufe": self.cufe,
            "cufe_corto": cufe_corto(self.cufe),
            "fecha_emision": fecha_iso(self.fecha_emision),
            "fecha_recepcion": self.fecha_recepcion.isoformat() if self.fecha_recepcion else None,
            "nit_emisor": self.nit_emisor,
            "nit_receptor": self.nit_receptor,
            "nombre_receptor": self.nombre_receptor,
            "consumidor_final": self.consumidor_final,
            "iva": dinero(self.iva),
            "inc": dinero(self.inc),
            "otros_impuestos": {k: dinero(v) for k, v in self.otros_impuestos.items() if v},
            "retenciones": {k: dinero(v) for k, v in self.retenciones.items() if v},
            "total": dinero(self.total),
            "base_derivada": dinero(self.base_derivada),
            "estado_dian": self.estado_dian,
            "advertencias": list(self.advertencias),
        }


@dataclass
class LecturaContenedor:
    nombre_archivo: str
    hoja: str
    fila_encabezado: int
    columnas: list[str]
    columnas_opcionales_faltantes: list[str]
    documentos: list[DocumentoDian]
    validaciones: list[dict]
    filas_vacias: int = 0

    @property
    def incluidos(self) -> list[DocumentoDian]:
        return [d for d in self.documentos if d.incluido]

    @property
    def excluidos(self) -> list[DocumentoDian]:
        return [d for d in self.documentos if not d.incluido]

    @property
    def prefijos_conocidos(self) -> set[str]:
        return {d.prefijo for d in self.documentos if d.prefijo}

    def nits_emisores(self) -> set[str]:
        return {d.nit_emisor for d in self.documentos if d.nit_emisor and texto_clave(d.grupo) == GRUPO_EMITIDO}

    def rango_fechas(self) -> tuple[Optional[date], Optional[date]]:
        fechas = [d.fecha_emision for d in self.documentos if d.fecha_emision and texto_clave(d.grupo) == GRUPO_EMITIDO]
        return (min(fechas), max(fechas)) if fechas else (None, None)

    def max_fecha_recepcion(self) -> Optional[datetime]:
        fechas = [d.fecha_recepcion for d in self.documentos if d.fecha_recepcion]
        return max(fechas) if fechas else None

    def resumen_tipos(self) -> list[dict]:
        grupos: dict[str, list[DocumentoDian]] = defaultdict(list)
        for d in self.documentos:
            grupos[d.tipo_documento].append(d)
        filas = []
        for tipo, docs in sorted(grupos.items(), key=lambda kv: -len(kv[1])):
            regla = regla_para_tipo(tipo)
            filas.append({
                "tipo_documento": tipo,
                "filas": len(docs),
                "total": dinero(sum((d.total for d in docs), CERO)),
                "prefijos": sorted({d.prefijo for d in docs if d.prefijo}),
                "incluido": bool(regla and regla.incluir),
                "tratamiento": regla.regla if regla else "Tipo no reconocido: queda en revisión y no suma",
                "categoria": regla.categoria if regla else "desconocido",
            })
        return filas

    def a_dict(self) -> dict:
        desde, hasta = self.rango_fechas()
        corte = self.max_fecha_recepcion()
        incluidos = self.incluidos
        return {
            "archivo": self.nombre_archivo,
            "hoja": self.hoja,
            "fila_encabezado": self.fila_encabezado,
            "filas": len(self.documentos),
            "columnas": self.columnas,
            "columnas_opcionales_faltantes": self.columnas_opcionales_faltantes,
            "nits_emisores": sorted(self.nits_emisores()),
            "fecha_desde": fecha_iso(desde),
            "fecha_hasta": fecha_iso(hasta),
            "corte_recepcion": corte.isoformat() if corte else None,
            "documentos_incluidos": len(incluidos),
            "documentos_excluidos": len(self.excluidos),
            "facturas": sum(1 for d in incluidos if d.categoria == "factura"),
            "notas_credito": sum(1 for d in incluidos if d.categoria == "nota_credito"),
            "notas_debito": sum(1 for d in incluidos if d.categoria == "nota_debito"),
            "neto": dinero(sum((d.total_firmado for d in incluidos), CERO)),
            "tipos": self.resumen_tipos(),
            "validaciones": self.validaciones,
        }


# ──────────────────────────────────────────────
# LECTURA DEL ARCHIVO
# ──────────────────────────────────────────────

def _demasiadas_filas(cantidad: int) -> ValueError:
    return ValueError(f"El contenedor tiene más de {MAX_FILAS} filas (al menos {cantidad}): exporta un período más corto.")


def _revisar_xlsx(contenido: bytes, tope: int) -> None:
    """
    Revisa el .xlsx antes de abrirlo con openpyxl, sin cargarlo en memoria.

    Un .xlsx es un ZIP: pocos MB pueden contener millones de filas idénticas.
    Si la hoja no declara <dimension>, openpyxl la recorre completa al abrirla,
    antes de que el tope de filas de _leer_filas pueda actuar. Por eso aquí:
      1. se rechaza el archivo si su tamaño descomprimido declarado supera
         MAX_XLSX_DESCOMPRIMIDO (zipfile nunca entrega más de lo declarado);
      2. se cuentan las filas de cada hoja descomprimiendo por bloques y se
         rechaza en cuanto superan el tope.
    """
    try:
        archivo = zipfile.ZipFile(io.BytesIO(contenido))
    except zipfile.BadZipFile:
        return  # openpyxl informará el error de formato
    with archivo:
        partes = archivo.infolist()
        if sum(p.file_size for p in partes) > MAX_XLSX_DESCOMPRIMIDO:
            raise ValueError(
                f"El contenedor ocupa más de {MAX_XLSX_DESCOMPRIMIDO // (1024 * 1024)} MB al descomprimirlo: "
                "exporta un período más corto."
            )
        for parte in partes:
            if not (parte.filename.startswith("xl/worksheets/") and parte.filename.endswith(".xml")):
                continue
            filas, cola = 0, b""
            with archivo.open(parte) as hoja:
                while bloque := hoja.read(1 << 20):
                    datos = cola + bloque
                    fin = 0
                    for coincidencia in _FILA_XML.finditer(datos):
                        filas += 1
                        fin = coincidencia.end()
                    if filas > tope:
                        raise _demasiadas_filas(filas)
                    # lo que sigue a la última fila contada puede ser el inicio de otra
                    cola = datos[max(fin, len(datos) - 64):]


def _leer_filas(contenido: bytes) -> tuple[str, list[tuple]]:
    """Lee la hoja del contenedor y se detiene en cuanto supera MAX_FILAS (no carga el resto)."""
    tope = MAX_FILAS + 15
    formato = _detect_excel_format(contenido, "contenedor DIAN")
    if formato == "xlsx":
        _revisar_xlsx(contenido, tope)
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        try:
            hoja = wb.worksheets[0]
            for ws in wb.worksheets:
                if any(texto_clave(c) == "cufe/cude"
                       for fila in ws.iter_rows(min_row=1, max_row=15, values_only=True) for c in fila or ()):
                    hoja = ws
                    break
            filas: list[tuple] = []
            for r in hoja.iter_rows(values_only=True):
                filas.append(tuple(r[:MAX_COLUMNAS]))
                if len(filas) > tope:
                    raise _demasiadas_filas(len(filas))
            return hoja.title, filas
        finally:
            wb.close()

    import xlrd

    libro = xlrd.open_workbook(file_contents=contenido, on_demand=True)
    hoja = libro.sheet_by_index(0)
    if hoja.nrows > tope:
        raise _demasiadas_filas(hoja.nrows)
    filas = []
    for i in range(hoja.nrows):
        fila = []
        for celda in hoja.row(i)[:MAX_COLUMNAS]:
            if celda.ctype == xlrd.XL_CELL_DATE:
                fila.append(xlrd.xldate_as_datetime(celda.value, libro.datemode))
            elif celda.ctype == xlrd.XL_CELL_EMPTY:
                fila.append(None)
            else:
                fila.append(celda.value)
        filas.append(tuple(fila))
    return hoja.name, filas


def _ubicar_encabezado(filas: list[tuple]) -> tuple[int, dict[str, int], list[str]]:
    objetivo = {texto_clave(c): c for c in TODAS_LAS_COLUMNAS}
    for i, fila in enumerate(filas[:15]):
        claves = [texto_clave(c) for c in fila]
        if "cufe/cude" in claves and "tipo de documento" in claves:
            indices: dict[str, int] = {}
            for j, clave in enumerate(claves):
                if clave in objetivo and objetivo[clave] not in indices:
                    indices[objetivo[clave]] = j
            columnas_archivo = [texto_limpio(c) for c in fila if texto_limpio(c)]
            return i, indices, columnas_archivo
    raise ValueError(
        "No se encontró el encabezado del contenedor DIAN: se esperaban las columnas "
        "«Tipo de documento» y «CUFE/CUDE» en las primeras 15 filas."
    )


def leer_contenedor(contenido: bytes, nombre_archivo: str = "contenedor.xlsx") -> LecturaContenedor:
    """Lee el Excel de documentos emitidos y aplica las reglas de inclusión."""
    hoja, filas = _leer_filas(contenido)
    idx_encabezado, indices, columnas_archivo = _ubicar_encabezado(filas)

    faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in indices]
    if faltantes:
        raise ValueError("Al contenedor DIAN le faltan columnas requeridas: " + ", ".join(faltantes))
    opcionales_faltantes = [c for c in TODAS_LAS_COLUMNAS if c not in indices and c not in COLUMNAS_REQUERIDAS]

    def valor(fila: tuple, columna: str):
        j = indices.get(columna)
        return fila[j] if j is not None and j < len(fila) else None

    prefijos_conocidos = {
        normalizar_prefijo(valor(f, "Prefijo")) for f in filas[idx_encabezado + 1:] if valor(f, "Prefijo")
    }
    prefijos_conocidos.discard("")

    documentos: list[DocumentoDian] = []
    vacias = 0
    importes_invalidos: list[str] = []
    importes_ambiguos: list[str] = []
    fechas_invalidas: list[int] = []

    for numero_fila, fila in enumerate(filas[idx_encabezado + 1:], start=idx_encabezado + 2):
        if not any(c not in (None, "") for c in fila):
            vacias += 1
            continue
        advertencias: list[str] = []
        tipo = texto_limpio(valor(fila, "Tipo de documento"))
        grupo = texto_limpio(valor(fila, "Grupo"))
        estado = texto_limpio(valor(fila, "Estado"))

        montos: dict[str, Decimal] = {}
        for col in ("IVA", "INC", "Total") + OTROS_IMPUESTOS + RETENCIONES:
            if col not in indices:
                continue
            bruto = valor(fila, col)
            monto, error = a_decimal_seguro(bruto)
            if importe_ambiguo(bruto):
                importes_ambiguos.append(f"fila {numero_fila}, {col}: «{texto_limpio(bruto)}» → {dinero(monto)}")
            if error:
                importes_invalidos.append(f"fila {numero_fila}, {col}")
                advertencias.append(f"Importe ilegible en «{col}»: se tomó como 0")
            montos[col] = monto

        fecha_emision = parsear_fecha(valor(fila, "Fecha Emisión"))
        if fecha_emision is None:
            fechas_invalidas.append(numero_fila)
            advertencias.append("Fecha Emisión ilegible")

        regla = regla_para_tipo(tipo)
        if texto_clave(grupo) != GRUPO_EMITIDO:
            categoria, signo, incluido = "excluido", 0, False
            motivo = f"Grupo «{grupo or 'vacío'}»: el módulo de ventas solo lee documentos emitidos"
        elif regla is None:
            categoria, signo, incluido = "desconocido", 0, False
            motivo = f"Tipo de documento no reconocido («{tipo}»): requiere una regla antes de sumarlo"
        elif not regla.incluir:
            categoria, signo, incluido, motivo = regla.categoria, 0, False, regla.regla
        elif texto_clave(estado) not in ESTADOS_APROBADOS:
            categoria, signo, incluido = "excluido", 0, False
            motivo = f"Estado DIAN «{estado or 'vacío'}»: solo se concilian documentos aprobados"
        else:
            categoria, signo, incluido, motivo = regla.categoria, regla.signo, True, regla.regla

        numero = numero_desde_dian(valor(fila, "Prefijo"), valor(fila, "Folio"), prefijos_conocidos)
        if incluido and numero is None:
            advertencias.append("Sin prefijo/folio legible: solo se podrá identificar por CUFE")

        documentos.append(DocumentoDian(
            fila=numero_fila,
            tipo_documento=tipo,
            categoria=categoria,
            signo=signo,
            incluido=incluido,
            motivo=motivo,
            cufe=normalizar_cufe(valor(fila, "CUFE/CUDE")),
            prefijo=normalizar_prefijo(valor(fila, "Prefijo")),
            folio=texto_limpio(valor(fila, "Folio")),
            numero=numero,
            fecha_emision=fecha_emision,
            fecha_recepcion=parsear_fecha_hora(valor(fila, "Fecha Recepción")),
            nit_emisor=normalizar_nit(valor(fila, "NIT Emisor")),
            nombre_emisor=texto_limpio(valor(fila, "Nombre Emisor")),
            nit_receptor=normalizar_nit(valor(fila, "NIT Receptor")),
            nombre_receptor=texto_limpio(valor(fila, "Nombre Receptor")),
            divisa=texto_limpio(valor(fila, "Divisa")),
            forma_pago=texto_limpio(valor(fila, "Forma de Pago")),
            medio_pago=texto_limpio(valor(fila, "Medio de Pago")),
            iva=montos.get("IVA", CERO),
            inc=montos.get("INC", CERO),
            # Solo valores distintos de cero: ahorra memoria con miles de filas.
            otros_impuestos={c: montos[c] for c in OTROS_IMPUESTOS if montos.get(c)},
            retenciones={c: montos[c] for c in RETENCIONES if montos.get(c)},
            total=montos.get("Total", CERO),
            estado_dian=estado,
            grupo=grupo,
            advertencias=advertencias,
        ))

    lectura = LecturaContenedor(
        nombre_archivo=nombre_archivo,
        hoja=hoja,
        fila_encabezado=idx_encabezado + 1,
        columnas=columnas_archivo,
        columnas_opcionales_faltantes=opcionales_faltantes,
        documentos=documentos,
        validaciones=[],
        filas_vacias=vacias,
    )
    lectura.validaciones = _validar(lectura, importes_invalidos, fechas_invalidas, importes_ambiguos)
    return lectura


# ──────────────────────────────────────────────
# VALIDACIONES DE ESTRUCTURA
# ──────────────────────────────────────────────

def _v(codigo: str, nivel: str, mensaje: str, detalle: str = "") -> dict:
    return {"codigo": codigo, "nivel": nivel, "mensaje": mensaje, "detalle": detalle}


def _validar(lectura: LecturaContenedor, importes_invalidos: list[str], fechas_invalidas: list[int],
             importes_ambiguos: list[str] | None = None) -> list[dict]:
    docs = lectura.documentos
    v: list[dict] = []
    total_cols = len([c for c in TODAS_LAS_COLUMNAS if c not in lectura.columnas_opcionales_faltantes])
    v.append(_v("columnas", "ok", "Columnas reconocidas con su nombre exacto",
                f"{total_cols} de {len(TODAS_LAS_COLUMNAS)}"))
    if lectura.columnas_opcionales_faltantes:
        v.append(_v("columnas_opcionales", "advertencia", "Faltan columnas opcionales",
                    ", ".join(lectura.columnas_opcionales_faltantes)))

    emitidos = sum(1 for d in docs if texto_clave(d.grupo) == GRUPO_EMITIDO)
    v.append(_v("grupo", "ok" if emitidos == len(docs) else "advertencia",
                "Grupo = Emitido en todas las filas" if emitidos == len(docs)
                else "Hay filas de otro grupo: se excluyen", f"{emitidos} de {len(docs)}"))

    nits = lectura.nits_emisores()
    nombres = {d.nombre_emisor for d in docs if d.nombre_emisor}
    v.append(_v("nit_emisor", "ok" if len(nits) == 1 else "error",
                "Un solo NIT Emisor" if len(nits) == 1 else "El contenedor mezcla varios NIT emisores",
                ", ".join(sorted(nits)) + (f" · {len(nombres)} variantes de nombre; se identifica por NIT" if len(nombres) > 1 else "")))

    cufes = [d.cufe for d in docs if d.cufe]
    repetidos = [c for c, n in Counter(cufes).items() if n > 1]
    largos = Counter(len(c) for c in cufes)
    v.append(_v("cufe", "ok" if not repetidos and len(cufes) == len(docs) else "advertencia",
                "CUFE/CUDE presente y sin repetidos" if not repetidos else "Hay CUFE/CUDE repetidos",
                f"{len(cufes)} con valor · longitudes {dict(largos)}" + (f" · repetidos: {len(repetidos)}" if repetidos else "")))

    claves = [d.numero.clave for d in lectura.incluidos if d.numero]
    claves_rep = [c for c, n in Counter(claves).items() if n > 1]
    v.append(_v("numero", "ok" if not claves_rep else "advertencia",
                "Prefijo + Folio sin repetidos" if not claves_rep else "Hay Prefijo + Folio repetidos",
                f"{len(lectura.prefijos_conocidos)} prefijos" + (f" · repetidos: {', '.join(claves_rep[:10])}" if claves_rep else "")))

    desde, hasta = lectura.rango_fechas()
    v.append(_v("fechas", "ok" if not fechas_invalidas else "advertencia",
                "Fecha Emisión legible" if not fechas_invalidas else "Hay fechas ilegibles",
                f"{fecha_iso(desde)} → {fecha_iso(hasta)}" + (f" · ilegibles en filas {fechas_invalidas[:10]}" if fechas_invalidas else "")))

    corte = lectura.max_fecha_recepcion()
    if corte and hasta:
        ultimo_dia = [d.fecha_recepcion for d in docs if d.fecha_emision == hasta and d.fecha_recepcion]
        ultima_hora = max(ultimo_dia) if ultimo_dia else corte
        if ultima_hora.hour < 22:
            v.append(_v("corte", "advertencia",
                        f"El {hasta.strftime('%d-%m-%Y')} puede llegar incompleto: el último documento recibido es de las {ultima_hora.strftime('%H:%M')}",
                        f"Fecha Recepción máxima {corte.isoformat(sep=' ')}"))

    v.append(_v("base", "info", "El contenedor no trae base ni subtotal: se deriva y se marca como derivada",
                "Total − IVA − INC − otros impuestos"))

    if importes_invalidos:
        v.append(_v("importes", "advertencia", "Hay importes ilegibles (se tomaron como 0)",
                    "; ".join(importes_invalidos[:10])))
    if importes_ambiguos:
        v.append(_v("importes_ambiguos", "advertencia",
                    "Hay importes en texto con un punto y 3 dígitos: se leyeron como miles (formato colombiano)",
                    "; ".join(importes_ambiguos[:10])))

    desconocidos = sorted({d.tipo_documento for d in docs if d.categoria == "desconocido"})
    if desconocidos:
        v.append(_v("tipos", "advertencia", "Hay tipos de documento sin regla: no suman hasta definirla",
                    ", ".join(desconocidos)))

    estados = Counter(d.estado_dian for d in docs)
    no_aprobados = {e: n for e, n in estados.items() if texto_clave(e) not in ESTADOS_APROBADOS}
    v.append(_v("estado", "ok" if not no_aprobados else "advertencia",
                "Todos los documentos están aprobados por la DIAN" if not no_aprobados
                else "Hay documentos no aprobados: se excluyen",
                " · ".join(f"{e or 'vacío'}: {n}" for e, n in estados.most_common())))
    return v
