"""
QA de punta a punta de FactuMatch Ventas (interfaz + API) con Playwright y datos sintéticos.

Levanta el servidor de QA (router real, Odoo y Claude simulados) y un servidor
estático del frontend, ejecuta los casos y deja resultados y capturas en
qa/ventas/salida/ (ignorada por git).

    pip install playwright && playwright install chromium     # una vez
    python qa/ventas/qa_ui.py

Variables: QA_PYTHON (intérprete con las dependencias del backend; por defecto el actual),
QA_SALIDA (carpeta de resultados).
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import time
import traceback
import urllib.error
import urllib.request

from playwright.sync_api import sync_playwright

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(os.path.dirname(AQUI))
FRONT = os.path.join(RAIZ, "frontend")
SALIDA = os.environ.get("QA_SALIDA", os.path.join(AQUI, "salida"))
PY = os.environ.get("QA_PYTHON", sys.executable)
API = "http://127.0.0.1:8765"
WEB = "http://127.0.0.1:8766"
TOKEN = "token-qa"

ADMIN = {"id": "qa-1", "name": "QA Admin", "email": "qa@example.com", "role": "Admin",
         "modules": "auditor,comparisons,checklist,reconciliation,purchasing,ventas"}
SIN_MODULO = {**ADMIN, "id": "qa-2", "name": "QA Sin Ventas", "role": "Operador",
              "modules": "auditor,comparisons,checklist,reconciliation,purchasing"}
LECTURA = {**ADMIN, "id": "qa-3", "name": "QA Lectura", "role": "Lectura"}


def http(metodo: str, ruta: str, datos: bytes | None = None, cabeceras: dict | None = None) -> tuple[int, bytes]:
    req = urllib.request.Request(API + ruta, data=datos, method=metodo, headers=cabeceras or {})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def esperar(url: str, segundos: int = 60) -> None:
    fin = time.time() + segundos
    while time.time() < fin:
        try:
            with urllib.request.urlopen(url, timeout=3) as r:
                if r.status == 200:
                    return
        except Exception:
            time.sleep(0.5)
    raise RuntimeError(f"No respondió {url}")


class Corrida:
    def __init__(self):
        self.casos: list[dict] = []
        self.consola: list[str] = []

    def caso(self, codigo: str, titulo: str, funcion, pagina=None) -> None:
        inicio = time.time()
        captura = None
        try:
            detalle = funcion() or ""
            resultado = "PASA"
        except Exception as exc:  # se registra y se sigue con los demás casos
            detalle = f"{type(exc).__name__}: {exc}"
            resultado = "FALLA"
            traceback.print_exc()
        if pagina is not None:
            captura = os.path.join(SALIDA, f"{codigo}.png")
            try:
                pagina.screenshot(path=captura, full_page=True)
            except Exception:
                captura = None
        self.casos.append({"codigo": codigo, "titulo": titulo, "resultado": resultado, "detalle": str(detalle)[:500],
                           "segundos": round(time.time() - inicio, 2), "captura": os.path.basename(captura) if captura else None})
        print(f"{resultado:5} {codigo} {titulo} — {str(detalle)[:140]}")


def nuevo_contexto(navegador, usuario: dict, con_odoo: bool = True, ancho: int = 1440):
    ctx = navegador.new_context(viewport={"width": ancho, "height": 900}, accept_downloads=True)
    ls = f"localStorage.setItem('factura_user', {json.dumps(json.dumps(usuario))});"
    if con_odoo:
        ls += f"localStorage.setItem('odoo_credentials', '{TOKEN}');"
    ctx.add_init_script(ls)
    app = open(os.path.join(FRONT, "js", "app.js"), encoding="utf-8").read().replace(
        'const API_URL = "https://factu-match.onrender.com";', f'const API_URL = "{API}";')
    ctx.route("**/js/app.js", lambda r: r.fulfill(status=200, content_type="application/javascript", body=app))
    ctx.route("**/script.google.com/**", lambda r: r.fulfill(status=200, content_type="application/json",
                                                             body='{"success":true,"data":[]}'))
    for externo in ("**/cdn.jsdelivr.net/**", "**/fonts.googleapis.com/**", "**/fonts.gstatic.com/**"):
        ctx.route(externo, lambda r: r.abort())
    return ctx


def abrir_ventas(ctx, corrida: Corrida | None = None):
    pg = ctx.new_page()
    if corrida is not None:
        pg.on("console", lambda m: corrida.consola.append(m.text) if m.type == "error" else None)
        pg.on("pageerror", lambda e: corrida.consola.append(f"pageerror: {e}"))
    pg.goto(f"{WEB}/index.html")
    pg.wait_for_timeout(600)
    pg.click("#nav-ventas")
    pg.wait_for_timeout(200)
    return pg


def cargar(pg, ruta: str, estado: str) -> None:
    pg.set_input_files("#vt-archivo", ruta)
    pg.wait_for_selector(f"#vt-estado-archivo:has-text('{estado}')", timeout=60000)


def conciliar(pg) -> None:
    pg.click("#btn-ventas-conciliar")
    pg.wait_for_selector("#vt-global", state="visible", timeout=120000)
    pg.wait_for_timeout(300)


def main() -> int:
    os.makedirs(SALIDA, exist_ok=True)
    sys.path.insert(0, AQUI)
    sys.path.insert(0, os.path.join(RAIZ, "backend"))
    from datos_sinteticos import construir, contenedor_dos_nit, fila_dian, crear_contenedor  # noqa: E402

    datos = construir()
    tmp = tempfile.mkdtemp(prefix="qa_ventas_")
    rutas = {
        "valido": os.path.join(tmp, "contenedor_sintetico.xlsx"),
        "dos_nit": os.path.join(tmp, "contenedor_dos_nit.xlsx"),
        "pdf": os.path.join(tmp, "no_es_excel.xlsx"),
        "xss": os.path.join(tmp, "contenedor_xss.xlsx"),
    }
    open(rutas["valido"], "wb").write(datos.contenedor)
    open(rutas["dos_nit"], "wb").write(contenedor_dos_nit())
    open(rutas["pdf"], "wb").write(b"%PDF-1.7 esto no es un Excel")
    open(rutas["xss"], "wb").write(crear_contenedor([fila_dian("FE", 1, 1000),
                                                     fila_dian("XS", 2, 10, tipo='<img src=x onerror="window.__xss=1">')]))

    procesos = [
        subprocess.Popen([PY, os.path.join(AQUI, "servidor_qa.py")], cwd=RAIZ,
                         stdout=subprocess.DEVNULL, stderr=open(os.path.join(SALIDA, "servidor_qa.log"), "w")),
        subprocess.Popen([sys.executable, "-m", "http.server", "8766", "--bind", "127.0.0.1"], cwd=FRONT,
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL),
    ]
    corrida = Corrida()
    try:
        esperar(API + "/api/ventas/estado-ia")
        esperar(WEB + "/index.html")
        with sync_playwright() as p:
            nav = p.chromium.launch()

            # ── Acceso y permisos ──
            ctx = nuevo_contexto(nav, ADMIN)
            pg = abrir_ventas(ctx, corrida)

            def qa01():
                assert pg.is_visible("#nav-ventas"), "el menú no aparece"
                assert pg.inner_text(".topbar-title") == "VENTAS DIAN VS ODOO"
                assert pg.is_visible("#view-ventas")
                assert all(pg.is_disabled(f"button.vt-tab[data-vt-tab='{t}']") for t in ("global", "detalle", "agente"))
                return "Menú visible, vista abierta y pestañas 2–4 bloqueadas hasta conciliar"
            corrida.caso("QA-01", "Acceso al módulo con permiso «ventas»", qa01, pg)

            ctx2 = nuevo_contexto(nav, SIN_MODULO)
            pg2 = ctx2.new_page()
            pg2.goto(f"{WEB}/index.html")
            pg2.wait_for_timeout(600)

            def qa02():
                assert not pg2.is_visible("#nav-ventas"), "el menú se ve sin permiso"
                pg2.evaluate("switchView('ventas')")
                assert not pg2.is_visible("#view-ventas"), "la vista abrió sin permiso"
                return "Sin el módulo: menú oculto y switchView('ventas') bloqueado"
            corrida.caso("QA-02", "Usuario sin el módulo no accede", qa02, pg2)
            ctx2.close()

            # ── Carga del contenedor ──
            def qa03():
                cargar(pg, rutas["pdf"], "ERROR")
                texto = pg.inner_text("#vt-validaciones")
                assert "formato Excel" in texto, texto
                assert pg.is_disabled("#btn-ventas-conciliar")
                return texto.strip()[:160]
            corrida.caso("QA-03", "Archivo que no es Excel", qa03, pg)

            def qa04():
                cargar(pg, rutas["dos_nit"], "CON ERRORES")
                assert "mezcla varios NIT" in pg.inner_text("#vt-validaciones")
                assert pg.is_disabled("#btn-ventas-conciliar")
                return "Validación bloqueante y botón deshabilitado"
            corrida.caso("QA-04", "Contenedor con dos NIT emisores", qa04, pg)

            def qa05():
                pg.fill("#vt-desde", "")
                pg.fill("#vt-hasta", "")
                cargar(pg, rutas["valido"], "VALIDADO")
                resumen = pg.inner_text(".vt-resumen-archivo")
                assert "268" in resumen and "12" in resumen, resumen
                assert pg.inner_text("#vt-resumen-columnas").startswith("32 de 32")
                assert pg.input_value("#vt-desde") == "2026-10-01" and pg.input_value("#vt-hasta") == "2026-10-15"
                filas_tipos = pg.locator("#vt-tabla-tipos tbody tr").count()
                assert filas_tipos == 7, filas_tipos
                assert "puede llegar incompleto" in pg.inner_text("#vt-validaciones")
                assert not pg.is_disabled("#btn-ventas-conciliar")
                return "268 incluidos, 12 excluidos, 7 tipos, 32/32 columnas, período sugerido y aviso de corte"
            corrida.caso("QA-05", "Contenedor válido: validaciones, tipos y columnas", qa05, pg)

            ctx3 = nuevo_contexto(nav, ADMIN, con_odoo=False)
            pg3 = abrir_ventas(ctx3)

            def qa06():
                cargar(pg3, rutas["valido"], "VALIDADO")
                assert "SIN CONECTAR" in pg3.inner_text("#vt-odoo-estado")
                assert pg3.is_disabled("#btn-ventas-conciliar")
                return "Sin credenciales de Odoo el botón queda deshabilitado"
            corrida.caso("QA-06", "Sin Odoo conectado", qa06, pg3)
            ctx3.close()

            # ── Conciliación global ──
            def qa07():
                conciliar(pg)
                kpis = pg.locator("#vt-kpis .metrica-val").all_inner_texts()
                assert "-822.610,40" in kpis[2], kpis
                assert kpis[3].strip().endswith("0,00"), kpis
                assert "CUADRA" in pg.inner_text("#vt-cuadra")
                no_ok = pg.locator("#vt-alcance .vt-alcance-item.no").count()
                assert no_ok == 0, f"{no_ok} criterios de alcance en rojo"
                return f"KPIs {kpis}; descomposición cuadra; 9 criterios de alcance en verde"
            corrida.caso("QA-07", "Conciliación y descomposición exacta", qa07, pg)

            def qa08():
                pg.click("[data-vt-componente='missing_in_dian']")
                pg.wait_for_selector("#vt-detalle", state="visible")
                activo = pg.inner_text("#vt-filtros .vt-filtro.active")
                filas = pg.locator("#vt-tabla-detalle tbody tr[data-vt-ref]").count()
                assert "missing_in_dian" in activo and filas == 9, (activo, filas)
                return f"Filtro «{activo}» con {filas} casos"
            corrida.caso("QA-08", "De la descomposición a los casos", qa08, pg)

            def qa09():
                pg.click("#vt-filtros [data-vt-filtro='todos']")
                info = pg.inner_text("#vt-pagina-info")
                assert info.startswith("Mostrando 1 a 50 de"), info
                pg.click("#btn-vt-siguiente")
                assert pg.inner_text("#vt-pagina-info").startswith("Mostrando 51 a 100"), pg.inner_text("#vt-pagina-info")
                pg.fill("#vt-buscar", "FE-22")
                pg.wait_for_timeout(400)
                filas = pg.locator("#vt-tabla-detalle tbody tr[data-vt-ref]")
                assert filas.count() == 1 and "Diferencia de valor" in filas.first.inner_text()
                return f"{info}; búsqueda FE-22 → 1 caso con diferencia de valor"
            corrida.caso("QA-09", "Filtros, paginación y búsqueda", qa09, pg)

            def qa10():
                pg.click("#vt-tabla-detalle tbody tr[data-vt-ref='dian:FE-22']")
                pg.wait_for_selector("#vt-evidencia .vt-lados", timeout=20000)
                texto = pg.inner_text("#vt-evidencia")
                assert "DIAN" in texto and "ODOO" in texto and "FE/2026/00022" in texto
                assert pg.locator("#vt-evidencia tr.vt-r-diferente").count() >= 1
                return "Evidencia lado a lado con la validación de total en rojo"
            corrida.caso("QA-10", "Evidencia de un documento", qa10, pg)

            def qa11():
                pg.fill("#vt-buscar", "")
                ruta_descarga = os.path.join(SALIDA, "reporte_qa.xlsx")
                pg.click("button.vt-tab[data-vt-tab='global']")
                with pg.expect_download(timeout=60000) as d:
                    pg.click("#btn-ventas-reporte")
                d.value.save_as(ruta_descarga)
                from openpyxl import load_workbook
                wb = load_workbook(ruta_descarga)
                assert len(wb.sheetnames) == 10, wb.sheetnames
                celdas = [c for ws in wb.worksheets for f in ws.iter_rows() for c in f
                          if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")]
                assert celdas and all(c.data_type == "s" for c in celdas), "fórmula inyectada sin neutralizar"
                return f"{d.value.suggested_filename}: {len(wb.sheetnames)} hojas; fórmula inyectada guardada como texto"
            corrida.caso("QA-11", "Reporte Excel", qa11, pg)

            # ── Agente ──
            def qa12():
                http("POST", "/qa/modo-ia/ok")
                pg.click("button.vt-tab[data-vt-tab='agente']")
                pg.wait_for_timeout(500)
                assert pg.inner_text("#vt-ia-estado") == "ACTIVO"
                pg.click("#btn-ventas-investigar")
                pg.wait_for_selector(".vt-hallazgo", timeout=120000)
                hallazgos = pg.locator(".vt-hallazgo")
                assert hallazgos.count() == 3, hallazgos.count()
                fe15 = hallazgos.nth(1).inner_text()
                assert "Causa Confirmada" in fe15, fe15[:200]           # búsqueda en vivo completa
                assert "Causa Probable" in hallazgos.nth(0).inner_text()
                pasos = pg.locator("#vt-traza .vt-paso").count()
                assert pasos == 4, pasos
                assert "IMPACTO CALCULADO POR REGLAS" in fe15
                return "3 hallazgos (hechos, hipótesis, acciones), 4 herramientas, FE-15 confirmada tras búsqueda en vivo"
            corrida.caso("QA-12", "Investigación de Claude con evidencia", qa12, pg)

            def qa13():
                pg.click("button.vt-tab[data-vt-tab='detalle']")
                pg.fill("#vt-buscar", "FE-15")
                pg.wait_for_timeout(400)
                pg.click("#vt-tabla-detalle tbody tr[data-vt-ref='dian:FE-15']")
                pg.wait_for_selector("[data-vt-investigar]", timeout=20000)
                pg.click("[data-vt-investigar]")
                assert pg.input_value("#vt-foco-tipo") == "documento" and pg.input_value("#vt-foco-ref") == "dian:FE-15"
                pg.click("#btn-ventas-investigar")
                pg.wait_for_selector(".vt-hallazgo", timeout=120000)
                return "Foco en dian:FE-15 desde la evidencia; investigación completada"
            corrida.caso("QA-13", "Investigar un documento desde su evidencia", qa13, pg)

            def qa14():
                http("POST", "/qa/modo-ia/falla")
                pg.select_option("#vt-foco-tipo", "global")
                pg.click("#btn-ventas-investigar")
                pg.wait_for_selector("#vt-ia-resultado .alerta", timeout=60000)
                texto = pg.inner_text("#vt-ia-resultado")
                assert "Análisis IA no disponible" in texto and "límite" in texto, texto[:200]
                pg.click("button.vt-tab[data-vt-tab='global']")
                with pg.expect_download(timeout=60000) as d:
                    pg.click("#btn-ventas-reporte")
                assert d.value.suggested_filename.endswith(".xlsx")
                http("POST", "/qa/modo-ia/ok")
                return "Falla del proveedor: aviso «no disponible» y el reporte se sigue descargando"
            corrida.caso("QA-14", "Falla de Claude sin afectar la conciliación", qa14, pg)

            def qa15():
                assert pg.evaluate("window.__xss") is None
                pg.click("button.vt-tab[data-vt-tab='carga']")
                cargar(pg, rutas["xss"], "VALIDADO")
                assert pg.evaluate("window.__xss") is None, "se ejecutó código inyectado"
                assert '<img src=x onerror="window.__xss=1">' in pg.inner_text("#vt-tabla-tipos")
                return "El texto malicioso del Excel se muestra literal y no se ejecuta"
            corrida.caso("QA-15", "Inyección de HTML desde el contenedor", qa15, pg)
            ctx.close()

            ctx4 = nuevo_contexto(nav, ADMIN)
            pg4 = abrir_ventas(ctx4)

            def qa16():
                http("POST", "/qa/modo-ia/sin_clave")
                cargar(pg4, rutas["valido"], "VALIDADO")
                conciliar(pg4)
                pg4.click("button.vt-tab[data-vt-tab='agente']")
                pg4.wait_for_timeout(600)
                assert pg4.inner_text("#vt-ia-estado") == "NO CONFIGURADO"
                assert pg4.is_disabled("#btn-ventas-investigar")
                assert "ANTHROPIC_API_KEY" in pg4.inner_text("#vt-ia-nota")
                http("POST", "/qa/modo-ia/ok")
                return "Sin clave: agente deshabilitado con explicación; conciliación disponible"
            corrida.caso("QA-16", "Sin ANTHROPIC_API_KEY", qa16, pg4)
            ctx4.close()

            ctx5 = nuevo_contexto(nav, LECTURA)
            pg5 = abrir_ventas(ctx5)

            def qa17():
                assert pg5.is_disabled("#vt-archivo") and pg5.is_disabled("#btn-ventas-conciliar")
                return "Rol Lectura: carga y conciliación deshabilitadas"
            corrida.caso("QA-17", "Rol Lectura", qa17, pg5)
            ctx5.close()

            ctx6 = nuevo_contexto(nav, ADMIN, ancho=1024)
            pg6 = abrir_ventas(ctx6)

            def qa18():
                cargar(pg6, rutas["valido"], "VALIDADO")
                conciliar(pg6)
                desborde = pg6.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth")
                assert desborde <= 1, f"desborde horizontal de {desborde}px"
                return "A 1024 px no hay desplazamiento horizontal de la página"
            corrida.caso("QA-18", "Pantalla de 1024 px", qa18, pg6)

            def qa19():
                pg6.click("#nav-comparisons")
                assert pg6.is_visible("#view-comparisons") and not pg6.is_visible("#view-ventas")
                pg6.click("#nav-dashboard")
                assert pg6.is_visible("#view-dashboard")
                pg6.click("#nav-ventas")
                assert pg6.is_visible("#view-ventas") and pg6.is_visible("#vt-global")
                return "Comparaciones y tablero siguen funcionando; al volver, ventas conserva el resultado"
            corrida.caso("QA-19", "Regresión de navegación con otros módulos", qa19, pg6)
            ctx6.close()

            # ── API ──
            def qa20():
                estado, cuerpo = http("GET", "/api/ventas/estado-ia")
                assert estado == 200 and b"clave-ficticia-de-qa" not in cuerpo
                assert http("GET", "/api/ventas/reporte?run_id=noexiste")[0] == 404
                assert http("GET", "/api/ventas/documento?run_id=noexiste&ref=dian:FE-1")[0] == 404
                limite = "----qa"
                cuerpo_form = (f"--{limite}\r\nContent-Disposition: form-data; name=\"credentials\"\r\n\r\nmalo\r\n"
                               f"--{limite}\r\nContent-Disposition: form-data; name=\"desde\"\r\n\r\n2026-10-01\r\n"
                               f"--{limite}\r\nContent-Disposition: form-data; name=\"hasta\"\r\n\r\n2026-10-15\r\n"
                               f"--{limite}\r\nContent-Disposition: form-data; name=\"dian\"; filename=\"c.xlsx\"\r\n"
                               f"Content-Type: application/octet-stream\r\n\r\n").encode() + datos.contenedor + f"\r\n--{limite}--\r\n".encode()
                estado, cuerpo = http("POST", "/api/ventas/conciliar", cuerpo_form, {"Content-Type": f"multipart/form-data; boundary={limite}"})
                assert estado == 400 and "Credenciales" in cuerpo.decode(), (estado, cuerpo[:200])
                return "estado-ia sin clave expuesta; 404 en conciliación inexistente; 400 con credenciales inválidas"
            corrida.caso("QA-20", "Contratos y errores de la API", qa20)

            def qa21():
                # «Failed to load resource» son registros de red: CDN bloqueados en QA y los 4xx esperados de
                # los casos negativos. Lo que no puede haber son errores de JavaScript de la aplicación.
                errores = [c for c in corrida.consola if not c.startswith("Failed to load resource")]
                assert not errores, errores[:5]
                return f"Sin errores de JavaScript ({len(corrida.consola)} avisos de red esperados)"
            corrida.caso("QA-21", "Consola del navegador sin errores de la aplicación", qa21)
            nav.close()
    finally:
        for proc in procesos:
            proc.terminate()

    resumen = {"pasan": sum(c["resultado"] == "PASA" for c in corrida.casos), "fallan": sum(c["resultado"] == "FALLA" for c in corrida.casos),
               "casos": corrida.casos, "fecha": time.strftime("%Y-%m-%d %H:%M:%S")}
    with open(os.path.join(SALIDA, "resultados_qa.json"), "w", encoding="utf-8") as f:
        json.dump(resumen, f, ensure_ascii=False, indent=2)
    print(f"\n{resumen['pasan']} pasan · {resumen['fallan']} fallan · resultados en {SALIDA}")
    return 0 if resumen["fallan"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
